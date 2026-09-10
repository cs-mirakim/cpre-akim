import openpyxl
import json
import re
import os
import base64
import fitz

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(ROOT_DIR, 'docs')
ASSETS_DIR = os.path.join(ROOT_DIR, 'assets', 'diagrams')
DATA_PATH = os.path.join(ROOT_DIR, 'data', 'app_data.json')

def get_b64(filename):
    path = os.path.join(ASSETS_DIR, filename)
    if os.path.exists(path):
        with open(path, 'rb') as f:
            return f"data:image/png;base64,{base64.b64encode(f.read()).decode('utf-8')}"
    return ""

img_q18 = get_b64('diagram_page_12_2.png')
img_q20 = get_b64('diagram_page_13_2.png')
img_q21 = get_b64('diagram_page_14_2.png')
img_q23 = get_b64('diagram_page_16_2.png')

# Load Excel Truth Keys
wb = openpyxl.load_workbook(os.path.join(DOCS_DIR, 'CorrectionAidForThePracticeExam_EN_2025-09-11.xlsx'), data_only=True)
ws = wb.active
col_idx = 2
excel_data = {}
letters = ['A', 'B', 'C', 'D', 'E', 'F']

while col_idx <= ws.max_column:
    val = ws.cell(row=1, column=col_idx).value
    if isinstance(val, int):
        q_num = val
        q_id = str(ws.cell(row=2, column=col_idx).value).strip()
        q_type = q_id[0]
        q_pts = float(ws.cell(row=4, column=col_idx).value)
        q_num_ans = ws.cell(row=5, column=col_idx).value
        
        answers = {}
        if q_type == 'K':
            num_rows = q_num_ans
            for i in range(num_rows):
                r = 7 + i
                key1 = ws.cell(row=r, column=col_idx+2).value
                answers[letters[i]] = True if key1 == 1 else False
            excel_data[q_num] = {
                'id': q_id, 'type': q_type, 'pts': q_pts, 'num_ans': q_num_ans, 'answers': answers
            }
            col_idx += 5
        else:
            correct_opts = []
            for i in range(6):
                r = 7 + i
                key = ws.cell(row=r, column=col_idx+1).value
                if key == 1:
                    correct_opts.append(letters[i])
            excel_data[q_num] = {
                'id': q_id, 'type': q_type, 'pts': q_pts, 'num_ans': q_num_ans, 'answers': correct_opts
            }
            col_idx += 3
    else:
        col_idx += 1

print(f"Loaded truth keys for {len(excel_data)} questions from Correction Aid.")

eus = [
    { "no": 1, "name": "Introduction and Overview of Requirements Engineering", "range": "Q1–Q3", "pts": 4, "qCount": 3 },
    { "no": 2, "name": "Fundamental Principles of Requirements Engineering", "range": "Q4–Q7", "pts": 6, "qCount": 4 },
    { "no": 3, "name": "Work Products and Documentation Practices", "range": "Q8–Q25", "pts": 30, "qCount": 18 },
    { "no": 4, "name": "Practices for Requirements Elaboration", "range": "Q26–Q35", "pts": 14, "qCount": 10 },
    { "no": 5, "name": "Process and Working Structure", "range": "Q36–Q37", "pts": 3, "qCount": 2 },
    { "no": 6, "name": "Management Practices for Requirements", "range": "Q38–Q43", "pts": 10, "qCount": 6 },
    { "no": 7, "name": "Tool Support", "range": "Q44–Q45", "pts": 3, "qCount": 2 }
]
